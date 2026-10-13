"""Add the audited bounded Scrapy extension to copies of current manuscripts.

These files remain drafts until every page is reviewed in Microsoft Word.
Never pool the two Scrapy bugs with independent projects or the mixed cohorts.
"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document

V2 = Path(__file__).resolve().parent.parent
CURRENT = V2 / "CURRENT_MANUSCRIPTS"
ANALYSIS = Path(__file__).resolve().parent / "scrapy_live_pair_v1/analysis_v1.json"
SOURCES = (
    ("ConveyorFlow_IEEE_Manuscript_v2_6_1_Current.docx",
     "ConveyorFlow_IEEE_Manuscript_v2_7_Exploratory_Draft_Rev7.docx", False, False),
    ("ConveyorFlow_AJSTR_Blinded_v2_6_1_Current.docx",
     "ConveyorFlow_AJSTR_Blinded_v2_7_Exploratory_Draft_Rev7.docx", True, True),
    ("ConveyorFlow_AJSTR_Unblinded_v2_6_1_Current.docx",
     "ConveyorFlow_AJSTR_Unblinded_v2_7_Exploratory_Draft_Rev7.docx", True, False),
)
FIGURE5 = Path(__file__).resolve().parent / "figures_rev6/Figure_5.png"

FROZEN_URL = (
    "https://github.com/sakanarm/ConveyorFlowExperimental/tree/"
    "691bc9e1397b45f9948069d4b6b44b0e48d8e826"
)


def windows_safe(path):
    value = str(Path(path).resolve())
    return "\\\\?\\" + value if len(value) >= 240 and not value.startswith("\\\\?\\") else value


def insert_before(anchor, text, style=None):
    paragraph = anchor.insert_paragraph_before(text, style=style)
    return paragraph


def find(doc, beginning):
    found = [paragraph for paragraph in doc.paragraphs
             if paragraph.text.startswith(beginning)]
    if len(found) != 1:
        raise ValueError("Expected exactly one anchor: " + beginning)
    return found[0]


def add_sentence(doc, beginning, sentence):
    find(doc, beginning).add_run(" " + sentence)


def replace_number_run(paragraph, old, new):
    matches = [run for run in paragraph.runs if run.text == old]
    if len(matches) != 1:
        raise ValueError(f"Expected one run {old!r} in {paragraph.text[:60]!r}")
    matches[0].text = new


def put_references_in_citation_order(doc, ajstr):
    kappa = find(doc, "We labeled 304 stage" if ajstr else
                 "Difficulty was operationalized over 304")
    statistics = find(doc, "A full run, not an individual task" if ajstr else
                      "The experimental unit is a complete run")
    replace_number_run(kappa, "20", "17")
    replace_number_run(statistics, "17", "18")
    replace_number_run(statistics, "19", "20")
    refs = {number: find(doc, f"[{number}] ") for number in range(17, 21)}
    refs[17]._element.addprevious(refs[20]._element)
    for old, new in ((20, 17), (17, 18), (18, 19), (19, 20)):
        replace_number_run(refs[old], str(old), str(new))


def replace_figure5(doc, ajstr):
    prefix = "Figure 5." if ajstr else "Fig. 5."
    captions = [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
    if len(captions) != 1:
        raise ValueError("Expected one Figure 5 caption")
    caption = captions[0]
    image_xml = caption._p.getprevious()
    image_paragraphs = [p for p in doc.paragraphs if p._p is image_xml]
    if len(image_paragraphs) != 1 or len(image_xml.xpath(".//wp:inline")) != 1:
        raise ValueError("Expected one image directly above Figure 5 caption")
    paragraph = image_paragraphs[0]
    old_inline = image_xml.xpath(".//wp:inline")[0]
    width = int(old_inline.xpath("./wp:extent")[0].get("cx"))
    paragraph.clear()
    shape = paragraph.add_run().add_picture(str(FIGURE5), width=width)
    shape._inline.docPr.set("descr", caption.text[:280])


def insert_table_before(doc, anchor, values):
    table = doc.add_table(rows=0, cols=len(values[0]))
    table.style = "Table Grid"
    for values_row in values:
        cells = table.add_row().cells
        for cell, value in zip(cells, values_row):
            cell.text = str(value)
    anchor._p.addprevious(table._tbl)
    return table


def build(doc, analysis, ajstr, blind):
    abstract = find(doc, "Abstract")
    abstract.text = abstract.text.replace(
        "and a container-verified repository repair.",
        "and a repository-repair attempt evaluated in a container.").replace(
        "The earlier six-block executable study is reported separately; the cohorts are not pooled.",
        "The earlier six-block study and a separate two-bug Scrapy follow-up "
        "are not pooled with this cohort. Each arm verified one of the two "
        "Scrapy repairs; no cross-project inference follows.")
    if ajstr:
        classical = find(doc, "Distributed task allocation has long")
        llm_work = find(doc, "LLM-agent frameworks such as")
    else:
        classical = find(doc, "Classical distributed task allocation")
        llm_work = find(doc, "LLM multi-agent systems such as")
    classical._element.addprevious(llm_work._element)
    put_references_in_citation_order(doc, ajstr)
    replace_figure5(doc, ajstr)
    if not ajstr:
        revision = find(doc, "Research revision 2.6.1")
        revision._element.getparent().remove(revision._element)
        contributions = find(doc, "Sakan Punyanon: investigation")
        contributions.text = contributions.text.replace(
            " The authors should confirm the final contribution wording before submission.", "")
        for heading in (
            "I. Frozen Ecological Real-LLM Allocation Main",
            "J. Supporting Executable ML Results",
            "K. Supporting Repository Repair Results",
            "L. Completed Repository Calibration",
            "M. Paired Ecological Main Results",
            "N. Post-Main Holdout Repository-Repair Feasibility",
        ):
            find(doc, heading).style = "Heading 2"
    if ajstr:
        method_anchor = find(doc, "3. Results and Discussion")
        insert_before(method_anchor, "2.6. Bounded Scrapy repository extension", "Heading 2")
        anchor = find(doc, "3.13. Interpretation")
        heading = "3.13. Bounded Scrapy repository extension"
        table_title = "Table 14. Bounded Scrapy repair extension"
        find(doc, "3.13. Interpretation").text = "3.14. Interpretation"
        find(doc, "3.14. Limitations").text = "3.15. Limitations"
    else:
        method_anchor = find(doc, "VI. RESULTS")
        insert_before(method_anchor, "K. Bounded Scrapy Repository Extension", "Heading 2")
        anchor = find(doc, "VII. DISCUSSION")
        heading = "P. Bounded Scrapy Repository Extension"
        table_title = "TABLE XV. BOUNDED SCRAPY REPAIR EXTENSION"
    insert_before(anchor, heading, "Heading 2")
    insert_before(method_anchor,
                  "We conducted a separate repair-only extension on the first two "
                  "previously unattempted Scrapy bugs in a pre-frozen public-source "
                  "queue. The buggy checkout failed its public test and the fixed "
                  "validator passed before any model call. The prompt contained only "
                  "the selected buggy production file and public failure trace. Five "
                  "non-visible public tests per case passed on both checkouts before "
                  "selection. Each case was presented on the same READY belt to "
                  "CF-Fit, Central-Rule-Matched, and Static-Owners, with one provider "
                  "request per arm, the same three-agent roster, the same prompt, "
                  "a 1,800-s horizon, and a 32,768-token output ceiling. Verification "
                  "required visible, regression, and fresh replay checks in Podman. "
                  "The task rank of 2 was a fixed exploratory allocation setting, "
                  "not an expert difficulty label.")
    insert_before(method_anchor,
                  "Restricted completion time assigns the 1,800-s horizon to a "
                  "repair not verified by the horizon; conditional completion "
                  "time is not imputed for failures. Busy fraction uses three "
                  "agents times the fixed horizon as its denominator. Gateway "
                  "response-cost values are reported only as provisional units, "
                  "not verified billed currency or monetary cost per verified job.")
    by_case = {case_id: {row["arm"]: row for row in analysis["rows"]
                         if row["case_id"] == case_id}
               for case_id in ("scrapy_39", "scrapy_34")}
    case_sentence = " ".join(
        f"For {case_id}, CF-Fit, rule-matched central allocation, and static "
        f"ownership verified {by_case[case_id]['CF_FIT']['verified']}/1, "
        f"{by_case[case_id]['CENTRAL_RULE_MATCHED']['verified']}/1, and "
        f"{by_case[case_id]['STATIC_OWNERS']['verified']}/1 repairs, respectively."
        for case_id in ("scrapy_39", "scrapy_34"))
    case_paragraph = insert_before(anchor, case_sentence)
    case_paragraph.add_run(" Table 14 summarizes these outcomes." if ajstr else
                           " Table XV summarizes these outcomes.")
    caption = insert_before(anchor, table_title)
    caption.paragraph_format.page_break_before = True
    caption.paragraph_format.keep_with_next = True
    arm_labels = (("CF_FIT", "CF-Fit"),
                  ("CENTRAL_RULE_MATCHED", "Central-Rule-Matched"),
                  ("STATIC_OWNERS", "Static-Owners"))
    values = [("Arm", "Verified", "Mean restricted time (s)",
               "Busy fraction (%)", "Output tokens")]
    for key, label in arm_labels:
        row = analysis["arms"][key]
        values.append((label, f"{row['verified_cases_of_two']}/2",
                       f"{row['mean_restricted_completion_seconds']:.1f}",
                       f"{100 * row['mean_busy_fraction_fixed_horizon']:.2f}",
                       str(row["total_output_tokens"])))
    insert_table_before(doc, anchor, values)
    reported = {key: sum(row["provider_cost_units_currency_unverified"]
                         for row in analysis["rows"] if row["arm"] == key)
                for key, _ in arm_labels}
    insert_before(anchor,
                  "Across both requests per arm, the gateway reported cost "
                  f"figures of {reported['CF_FIT']:.6f}, "
                  f"{reported['CENTRAL_RULE_MATCHED']:.6f}, and "
                  f"{reported['STATIC_OWNERS']:.6f} units for CF-Fit, "
                  "Central-Rule-Matched, and Static-Owners, respectively. "
                  "These are provisional gateway estimates; the MFEC tariff, "
                  "currency, and invoice have not been reconciled. Thus, they "
                  "are not verified billed costs or a monetary efficiency ranking.")
    insert_before(anchor,
                  "The two bugs belong to one "
                  "repository project and all arms used the selected A1 model; "
                  "provider calls are not independent project-level samples. "
                  "This extension does not add a full ML pipeline, establish "
                  "cross-project generalization, or isolate the effect of "
                  "stand-down. The earlier six- and eleven-block executable "
                  "cohorts remain separate and are not pooled with these cases.")
    future_anchor = find(doc, "Data and code availability" if ajstr
                         else "DATA AND CODE AVAILABILITY")
    insert_before(future_anchor, "Future work", "Heading 2")
    insert_before(future_anchor,
                  "A preregistered extension should sample independent public "
                  "repository projects and ML corpora, size the project-level "
                  "paired comparison for a stated precision target, and validate "
                  "task difficulty and Ability Ranks with independent human "
                  "raters before observing outcomes. It should compare matched "
                  "homogeneous and heterogeneous teams and isolate fit and "
                  "stand-down by ablation. Belt-order sensitivity needs two "
                  "separate manipulations: the order in which READY tasks are "
                  "exposed and the order in which agents reach or claim them. "
                  "With workloads and decision rules held fixed, future runs "
                  "should measure collisions, waiting, starvation, completion, "
                  "utilization, and cost under these orderings. Controlled "
                  "coordinator-only, shared claim-store, and provider outages "
                  "should delimit the failure-domain claim. Finally, the "
                  "gateway's dated tariff, usage records, and invoice must be "
                  "reconciled before reporting billed cost per verified job. "
                  "These are proposed tests, not results of the present study.")
    if ajstr:
        for number, sentence in (
            (12, "Table 12 summarizes the separate 11-block executable cohort."),
            (13, "Table 13 reports the matched decision-layer fault outcomes."),
        ):
            insert_before(find(doc, f"Table {number}."), sentence)
        image_only_headings = [p for p in doc.paragraphs
                               if p.style.name.startswith("Heading")
                               and not p.text.strip()
                               and p._p.xpath(".//w:drawing")]
        if len(image_only_headings) != 1:
            raise ValueError("Expected Figure 16 image-only heading")
        image_only_headings[0].style = "Normal"
        threat = find(doc, "Frozen simulation seeds, source hashes,")
        threat.text = threat.text.replace(
            "and component-failure tests remain open.",
            "and production-scale component-failure tests remain open.")
    else:
        add_sentence(doc, "Static controls are fixed skill ownership",
                     "Table I summarizes the prespecified controls and study cells.")
        add_sentence(doc, "Against the static controls, CF-Fit increased",
                     "Figure 6 shows the paired cost contrasts and uncertainty bands.")
        add_sentence(doc, "The radar chart is a visualization aid only.",
                     "Figure 7 shows this descriptive profile.")
        add_sentence(doc, "The large P95 reduction also has a competing-risk explanation:",
                     "Figure 8 displays the verified-throughput contrast.")
        add_sentence(doc, "Assessment increased cost by 4.6",
                     "Table IV and Figure 9 present the component contrasts.")
        add_sentence(doc, "The fallback analysis combines mixed",
                     "Table V defines the reported fallback comparisons.")
        add_sentence(doc, "Relative to F0/F1, F2 spent longer",
                     "Figure 10 shows the dead-letter contrasts.")
        add_sentence(doc, "The cost-throughput Pareto frontier contains",
                     "Figures 11 and 12 display the Pareto and competing-risk views, respectively.")
        add_sentence(doc, "A separate pre-Main calibration called",
                     "Figure 13 shows the descriptive pass-rate profile.")
        add_sentence(doc, "The original Real-LLM microtask study used",
                     "Figure 14 shows the descriptive policy profile.")
        add_sentence(doc, "A separately frozen extension added",
                     "Figure 15 shows the supplementary descriptive profiles.")
        add_sentence(doc, "Full-minus-ablation effects identify",
                     "Table IV reports full-minus-ablation results.")
        insert_before(find(doc, "TABLE XII."),
                      "Table XII summarizes the audited six-block ecological main.")
        insert_before(find(doc, "TABLE XIII."),
                      "Table XIII summarizes the separate 11-block executable cohort.")
        insert_before(find(doc, "TABLE XIV."),
                      "Table XIV reports the matched decision-layer fault outcomes.")
        robustness = find(doc, "Resource-mapping effects were recomputed")
        robustness.text = robustness.text.replace(
            "was stable across all three mappings=True",
            "was stable across all three mappings")
        no_standdown = find(doc, "Removing stand-down changed completion")
        no_standdown.text = no_standdown.text.replace(
            "-0.00 percentage points", "0.00 percentage points").replace(
            "After Holm correction, significant outcomes were: none of the reported outcomes.",
            "No reported outcome remained significant after Holm correction.")
        threat = find(doc, "Internal validity is strengthened by frozen seeds")
        threat.text = threat.text.replace(
            "and genuine component-failure tests remain future work.",
            "and production-scale component-failure tests remain future work.")
    conclusion = find(doc, "The independently audited executable cohort demonstrates")
    conclusion.add_run(
        " A separate two-bug Scrapy follow-up verified one repair per arm within "
        "one repository; it is not evidence of a policy advantage or "
        "cross-project reliability.")
    if blind:
        availability = (
            "The replication package contains experimental code, frozen "
            "configurations, public-data manifests, analysis scripts, derived "
            "results, and editable diagrams. Its author-identifying GitHub URL "
            "is omitted from this blinded manuscript and will be supplied to "
            "the editorial office separately. Public datasets and bug "
            "benchmarks remain at their upstream sources under their licenses. "
            "API credentials, raw provider responses, and heavy container "
            "artifacts are not public. Repository hashes and audit capsules "
            "therefore do not by themselves recreate historical provider calls; "
            "a fresh run needs independent provider access and records new "
            "dated outputs."
        )
    else:
        availability = (
            "Experimental code, frozen configurations, public-data manifests, "
            "analysis scripts, derived aggregate results, and editable diagrams "
            f"are available at {FROZEN_URL} (frozen commit 691bc9e). "
            "The v2.7 extension is in major_revision_v2_7, with "
            "REPRODUCE_V27.ps1 documenting the checks and optional fresh run. "
            "Public datasets and bug benchmarks remain at their upstream "
            "sources under their licenses. API credentials, raw provider "
            "responses, heavy container artifacts, and working manuscripts are "
            "not public. Repository hashes and audit capsules do not by "
            "themselves recreate historical provider calls; a fresh run needs "
            "independent provider access and records its own model identity, "
            "timestamps, and outputs."
        )
    if ajstr:
        find(doc, "Code, frozen configurations,").text = availability
    else:
        find(doc, "The public repository provides experimental source,").text = availability
        duplicate = find(doc, "Code and editable diagrams are available at")
        duplicate._element.getparent().remove(duplicate._element)
    doc.core_properties.subject = "Audited bounded Scrapy extension, separate from main cohorts"
    if blind:
        doc.core_properties.author = ""
        doc.core_properties.last_modified_by = ""


def main():
    analysis = json.loads(ANALYSIS.read_text(encoding="utf-8"))
    if (analysis["status"] != "BOUNDED_DESCRIPTIVE_ONE_REPOSITORY_NOT_MAIN" or
            analysis["cases"] != 2 or
            analysis["independent_repository_projects"] != 1):
        raise ValueError("Audited bounded analysis required")
    for source_name, target_name, ajstr, blind in SOURCES:
        target = CURRENT / target_name
        if target.exists():
            raise FileExistsError("Refuse to overwrite " + str(target))
        doc = Document(windows_safe(CURRENT / source_name))
        before = (len(doc.inline_shapes), len(doc.element.body.xpath('.//m:oMath')))
        build(doc, analysis, ajstr, blind)
        after = (len(doc.inline_shapes), len(doc.element.body.xpath('.//m:oMath')))
        if before != after:
            raise ValueError("Figure/native equation count changed")
        doc.save(windows_safe(target))
        print(target_name, "figures", after[0], "equations", after[1],
              "tables", len(doc.tables), flush=True)


if __name__ == "__main__":
    main()
