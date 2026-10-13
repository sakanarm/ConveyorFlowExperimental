"""Create citation-verified Rev9 copies; never overwrite earlier manuscripts."""

from __future__ import annotations

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parent.parent
CURRENT = ROOT / "CURRENT_MANUSCRIPTS"
NAMES = (
    ("ConveyorFlow_IEEE_Manuscript_v2_7_Exploratory_Draft_Rev7.docx", False),
    ("ConveyorFlow_AJSTR_Blinded_v2_7_Exploratory_Draft_Rev7.docx", True),
    ("ConveyorFlow_AJSTR_Unblinded_v2_7_Exploratory_Draft_Rev7.docx", True),
)


def windows_safe(path: Path) -> str:
    value = str(path.resolve())
    return "\\\\?\\" + value if len(value) >= 240 and not value.startswith("\\\\?\\") else value


def one(doc, prefix: str):
    hits = [paragraph for paragraph in doc.paragraphs if paragraph.text.startswith(prefix)]
    if len(hits) != 1:
        raise ValueError(f"Expected one paragraph beginning {prefix!r}; found {len(hits)}")
    return hits[0]


def change_reference(doc, number: int, value: str):
    one(doc, f"[{number}] ").text = value


def insert_before(anchor, value: str):
    return anchor.insert_paragraph_before(value)


AJSTR_ADDITIONS = (
    "[22] Abuhaimed, S.; Sen, S. Effective Task Allocation in Ad Hoc Human-Agent Teams. In HHAI2022: Augmenting Human Intellect, 2022; pp 171–183. https://doi.org/10.3233/FAIA220197.",
    "[23] Yue, Y.; Zhang, G.; Liu, B.; Wan, G.; Wang, K.; Cheng, D.; Qi, Y. MasRouter: Learning to Route LLMs for Multi-Agent Systems. In Proceedings of ACL, 2025; pp 15549–15572. https://doi.org/10.18653/v1/2025.acl-long.757.",
    "[24] Zhou, H.; Geng, H.; Xue, X.; Kang, L.; Qin, Y.; Wang, Z.; Yin, Z.; Bai, L. ReSo: A Reward-driven Self-organizing LLM-based Multi-Agent System for Reasoning Tasks. In Proceedings of EMNLP, 2025; pp 15979–15998. https://doi.org/10.18653/v1/2025.emnlp-main.808.",
    "[25] Huang, Q.; Vora, J.; Liang, P.; Leskovec, J. MLAgentBench: Evaluating Language Agents on Machine Learning Experimentation. In Proceedings of ICML; PMLR, 2024; Vol. 235, pp 20271–20309. https://proceedings.mlr.press/v235/huang24y.html.",
    "[26] Jimenez, C. E.; Yang, J.; Wettig, A.; Yao, S.; Pei, K.; Press, O.; Narasimhan, K. SWE-bench: Can Language Models Resolve Real-world Github Issues? In Proceedings of ICLR, 2024. https://proceedings.iclr.cc/paper_files/paper/2024/hash/edac78c3e300629acfe6cbe9ca88fb84-Abstract-Conference.html.",
)

IEEE_ADDITIONS = (
    "[21] R. Widyasari et al., “BugsInPy: A database of existing bugs in Python programs to enable controlled testing and debugging studies,” in Proc. ESEC/FSE, pp. 1556–1560, 2020, doi: 10.1145/3368089.3417943.",
    "[22] S. Abuhaimed and S. Sen, “Effective task allocation in ad hoc human-agent teams,” in Proc. HHAI2022, pp. 171–183, 2022, doi: 10.3233/FAIA220197.",
    "[23] Y. Yue et al., “MasRouter: Learning to route LLMs for multi-agent systems,” in Proc. ACL, pp. 15549–15572, 2025, doi: 10.18653/v1/2025.acl-long.757.",
    "[24] H. Zhou et al., “ReSo: A reward-driven self-organizing LLM-based multi-agent system for reasoning tasks,” in Proc. EMNLP, pp. 15979–15998, 2025, doi: 10.18653/v1/2025.emnlp-main.808.",
    "[25] Q. Huang, J. Vora, P. Liang, and J. Leskovec, “MLAgentBench: Evaluating language agents on machine learning experimentation,” in Proc. ICML, PMLR, vol. 235, pp. 20271–20309, 2024.",
    "[26] C. E. Jimenez et al., “SWE-bench: Can language models resolve real-world GitHub issues?,” in Proc. ICLR, 2024.",
)


def revise(doc, ajstr: bool):
    if ajstr:
        change_reference(doc, 1,
            "[1] Wu, Q.; Bansal, G.; Zhang, J.; Wu, Y.; Li, B.; Zhu, E.; Jiang, L.; Zhang, X.; Zhang, S.; Liu, J.; et al. AutoGen: Enabling Next-Gen LLM Applications via Multi-Agent Conversations. In Proceedings of COLM, 2024. https://openreview.net/pdf?id=BAakY1hNKS.")
        change_reference(doc, 2,
            "[2] Hong, S.; Zhuge, M.; Chen, J.; Zheng, X.; Cheng, Y.; Wang, J.; Zhang, C.; Wang, Z.; Yau, S. K. S.; Lin, Z.; et al. MetaGPT: Meta Programming for a Multi-Agent Collaborative Framework. In Proceedings of ICLR, 2024. https://proceedings.iclr.cc/paper_files/paper/2024/hash/6507b115562bb0a305f1958ccc87355a-Abstract-Conference.html.")
        change_reference(doc, 3,
            "[3] Qian, C.; Liu, W.; Liu, H.; Chen, N.; Dang, Y.; Li, J.; Yang, C.; Chen, W.; Su, Y.; Cong, X.; et al. ChatDev: Communicative Agents for Software Development. In Proceedings of ACL, 2024; pp 15174–15186. https://doi.org/10.18653/v1/2024.acl-long.810.")
        change_reference(doc, 5,
            "[5] Chen, W.; Su, Y.; Zuo, J.; Yang, C.; Yuan, C.; Chan, C.-M.; Yu, H.; Lu, Y.; Hung, Y.-H.; Qian, C.; et al. AgentVerse: Facilitating Multi-Agent Collaboration and Exploring Emergent Behaviors. In Proceedings of ICLR, 2024. https://proceedings.iclr.cc/paper_files/paper/2024/hash/578e65cdee35d00c708d4c64bce32971-Abstract-Conference.html.")
        change_reference(doc, 6,
            "[6] Guo, T.; Chen, X.; Wang, Y.; Chang, R.; Pei, S.; Chawla, N. V.; Wiest, O.; Zhang, X. Large Language Model Based Multi-Agents: A Survey of Progress and Challenges. In Proceedings of IJCAI, 2024; pp 8048–8057. https://doi.org/10.24963/ijcai.2024/890.")
        change_reference(doc, 11,
            "[11] Liu, Y.; Iter, D.; Xu, Y.; Wang, S.; Xu, R.; Zhu, C. G-Eval: NLG Evaluation Using GPT-4 with Better Human Alignment. In Proceedings of EMNLP, 2023; pp 2511–2522. https://doi.org/10.18653/v1/2023.emnlp-main.153.")
        change_reference(doc, 16,
            "[16] Lu, S.; Guo, D.; Ren, S.; Huang, J.; Svyatkovskiy, A.; Blanco, A.; Clement, C.; Drain, D.; Jiang, D.; Tang, D.; et al. CodeXGLUE: A Machine Learning Benchmark Dataset for Code Understanding and Generation. In NeurIPS Datasets and Benchmarks, 2021. https://datasets-benchmarks-proceedings.neurips.cc/paper/2021/hash/c16a5320fa475530d9583c34fd356ef5-Abstract-round1.html.")
        change_reference(doc, 17,
            "[17] Cohen, J. Weighted Kappa: Nominal Scale Agreement with Provision for Scaled Disagreement or Partial Credit. Psychol. Bull. 1968, 70 (4), 213–220. https://doi.org/10.1037/h0026256.")
        change_reference(doc, 18,
            "[18] Wilcoxon, F. Individual Comparisons by Ranking Methods. Biometrics Bull. 1945, 1 (6), 80–83. https://doi.org/10.2307/3001968.")
        change_reference(doc, 19,
            "[19] Efron, B. Bootstrap Methods: Another Look at the Jackknife. Ann. Stat. 1979, 7 (1), 1–26. https://doi.org/10.1214/aos/1176344552.")
        anchor = one(doc, "3.15. Limitations")
        additions = AJSTR_ADDITIONS
    else:
        change_reference(doc, 1,
            "[1] Q. Wu et al., “AutoGen: Enabling next-gen LLM applications via multi-agent conversations,” in Proc. COLM, 2024.")
        change_reference(doc, 16,
            "[16] S. Lu et al., “CodeXGLUE: A machine learning benchmark dataset for code understanding and generation,” in NeurIPS Datasets and Benchmarks, 2021.")
        change_reference(doc, 17,
            "[17] J. Cohen, “Weighted kappa: Nominal scale agreement with provision for scaled disagreement or partial credit,” Psychol. Bull., vol. 70, no. 4, pp. 213–220, 1968, doi: 10.1037/h0026256.")
        change_reference(doc, 18,
            "[18] F. Wilcoxon, “Individual comparisons by ranking methods,” Biometrics Bull., vol. 1, no. 6, pp. 80–83, 1945, doi: 10.2307/3001968.")
        change_reference(doc, 19,
            "[19] B. Efron, “Bootstrap methods: Another look at the jackknife,” Ann. Stat., vol. 7, no. 1, pp. 1–26, 1979, doi: 10.1214/aos/1176344552.")
        first_bugsinpy = one(doc, "Six mixed-workload blocks each admitted")
        first_bugsinpy.text = first_bugsinpy.text.replace("one BugsInPy repair at", "one BugsInPy repair [21] at", 1)
        anchor = one(doc, "VIII. THREATS TO VALIDITY")
        additions = IEEE_ADDITIONS

    # The new papers are discussed only after all earlier references have appeared.
    insert_before(anchor,
        "Related allocation studies delimit the comparison. Abuhaimed and Sen varied "
        "which member allocated tasks in human-agent teams [22]; this makes the "
        "decision maker relevant but does not establish a result for LLM teams. "
        "MasRouter uses a cascaded controller for collaboration mode, role, and "
        "model routing [23], whereas ReSo uses reward-driven agent selection for "
        "reasoning tasks [24]. Neither is a rule-matched test of local volunteering "
        "and temporary stand-down on a shared READY belt. Their reported gains "
        "should therefore not be read as policy-level benchmarks for our workloads.")
    insert_before(anchor,
        "MLAgentBench evaluates agents that conduct machine-learning experiments "
        "in executable workspaces [25], and SWE-bench evaluates patches for "
        "real GitHub issues across repositories [26]. These benchmarks show why "
        "our limited number of reused ML corpora and repository projects cannot "
        "support a broad claim about ML research automation or general "
        "repository repair. We cite them as external validity context, not as "
        "experimental baselines or additional ConveyorFlow results.")
    last_existing = one(doc, "[21] ") if ajstr else one(doc, "[20] ")
    for entry in additions:
        added = doc.add_paragraph(entry, style=last_existing.style)
        last_existing._p.addnext(added._p)
        last_existing = added


def main():
    for source, ajstr in NAMES:
        destination = CURRENT / source.replace("Rev7.docx", "Rev9_ReferencesVerified.docx")
        if destination.exists():
            raise FileExistsError(destination)
        doc = Document(windows_safe(CURRENT / source))
        figures_before = len(doc.inline_shapes)
        equations_before = len(doc.element.body.xpath(".//m:oMath"))
        tables_before = len(doc.tables)
        revise(doc, ajstr)
        if (len(doc.inline_shapes), len(doc.element.body.xpath(".//m:oMath")), len(doc.tables)) != (
            figures_before, equations_before, tables_before
        ):
            raise ValueError("Figures, equations, or tables changed")
        doc.save(windows_safe(destination))
        print(destination.name, "figures", figures_before,
              "equations", equations_before, "tables", tables_before)


if __name__ == "__main__":
    main()
