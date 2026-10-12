"""Structural and claim-scope QA for the three v2.6.1 Word drafts.

This does not replace page-by-page visual review in Microsoft Word.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import re
import zipfile

from docx import Document

from build_current_manuscripts_v2_6_1 import CURRENT, SOURCES, windows_safe


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def media(path: Path) -> dict:
    with zipfile.ZipFile(windows_safe(path)) as archive:
        return {name: sha(archive.read(name)) for name in archive.namelist()
                if name.startswith("word/media/")}


def audit() -> list[dict]:
    rows = []
    for source_name, output_name, ajstr, blind in SOURCES:
        source, output = CURRENT / source_name, CURRENT / output_name
        old, new = Document(windows_safe(source)), Document(windows_safe(output))
        text = "\n".join(p.text for p in new.paragraphs)
        expected = (
            "22,500-run simulation", "11 paired blocks", "17/22", "14/22",
            "18/22", "shared belt/claim-store outage", "earlier six-block",
            "not pooled", "previously exposed FastAPI", "32,768-token",
            "unknown cost", "not a universal",
        )
        missing = [value for value in expected if value.lower() not in text.lower()]
        if missing:
            raise ValueError(f"{output_name}: missing scope/data {missing}")
        if media(source) != media(output):
            raise ValueError(f"{output_name}: embedded figures/media changed")
        original = (len(old.inline_shapes), len(old.element.body.xpath('.//m:oMath')),
                    len(old.tables))
        current = (len(new.inline_shapes), len(new.element.body.xpath('.//m:oMath')),
                   len(new.tables))
        if current != (original[0], original[1], original[2] + 2):
            raise ValueError(f"{output_name}: figure/equation/table drift {original} -> {current}")
        if ajstr:
            abstract = next(p.text for p in new.paragraphs if p.text.startswith("Abstract:"))
            words = len(re.findall(r"\b[\w-]+\b", abstract))
            if words > 250:
                raise ValueError(f"{output_name}: AJSTR abstract too long: {words}")
            for heading in ("3.12. Separate 11-block executable cohort and fault results",
                            "3.13. Interpretation", "3.14. Limitations"):
                if text.count(heading) != 1:
                    raise ValueError(f"{output_name}: heading count drift: {heading}")
        else:
            words = len(re.findall(r"\b[\w-]+\b", next(
                p.text for p in new.paragraphs if p.text.startswith("Abstract"))))
        if blind:
            with zipfile.ZipFile(windows_safe(output)) as archive:
                xml = b"\n".join(archive.read(name) for name in archive.namelist()
                                 if name.endswith(".xml") or name.endswith(".rels"))
            lowered = xml.lower()
            for identifier in (b"sakan", b"punyanon", b"chaiyaporn",
                               b"aomduan", b"68140010", b"github.com/sakanarm"):
                if identifier in lowered:
                    raise ValueError(f"Blind identifier leaked: {identifier!r}")
        rows.append({"file": output_name, "source": source_name,
                     "sha256": sha(output.read_bytes()), "figures": current[0],
                     "native_equations": current[1], "tables": current[2],
                     "abstract_words": words,
                     "structural_and_claim_scope": "passed",
                     "visual_qa": "not_performed_word_required"})
    return rows


if __name__ == "__main__":
    print(json.dumps(audit(), indent=2))
