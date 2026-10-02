from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PRESENTATION = ROOT / "presentation"
TALK = PRESENTATION / "ConveyorFlow_Advisor_Talk_10min_TH.md"
QA = PRESENTATION / "ConveyorFlow_Advisor_QA_TH.md"
OUTPUT = PRESENTATION / "ConveyorFlow_Advisor_Talk_and_QA_TH.docx"

FONT = "TH Sarabun New"
NAVY = "17324D"
TEAL = "087783"
ORANGE = "DB7B2B"
LIGHT_BLUE = "EAF3F5"
LIGHT_ORANGE = "F9E9DB"
GRAY = "5A6875"


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_number(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run, 10, color=GRAY)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])


def add_header(section) -> None:
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("ConveyorFlow  |  Advisor Review v5")
    set_run_font(run, 10, bold=True, color=TEAL)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    add_header(section)
    add_page_number(section)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(15)
    normal.font.color.rgb = RGBColor.from_string("1F2933")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05

    for name, size, color in (
        ("Title", 30, NAVY),
        ("Heading 1", 24, NAVY),
        ("Heading 2", 19, TEAL),
        ("Heading 3", 16, ORANGE),
    ):
        style = document.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(4)


def add_markdown(document: Document, path: Path, skip_title: bool = True) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    for raw in lines:
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            if not skip_title:
                document.add_heading(line[2:].strip(), level=1)
            continue
        if line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
            continue
        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
            continue
        if line.startswith("> "):
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.45)
            p.paragraph_format.right_indent = Cm(0.35)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            shade_paragraph(p, LIGHT_BLUE)
            run = p.add_run(line[2:].strip())
            set_run_font(run, 14, bold=True, color=TEAL)
            continue
        if line.startswith("- "):
            p = document.add_paragraph(style="List Bullet")
            run = p.add_run(line[2:].strip())
            set_run_font(run, 14)
            continue
        p = document.add_paragraph()
        run = p.add_run(line)
        set_run_font(run, 15)


def add_cover(document: Document) -> None:
    document.add_paragraph("\n\n")
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ConveyorFlow")
    set_run_font(r, 34, bold=True, color=NAVY)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("บทพูดนำเสนอ 10 นาทีและคำถาม–คำตอบสำหรับอาจารย์")
    set_run_font(r, 23, bold=True, color=TEAL)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ฉบับสอดคล้องกับ Advisor Presentation v5\nSimulation 22,500 runs + Real-LLM Main and Extension")
    set_run_font(r, 15, color=GRAY)

    table = document.add_table(rows=3, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(5.0)
    table.columns[1].width = Cm(10.0)
    rows = (
        ("การใช้งาน", "ซ้อมนำเสนอ 10 นาที และใช้ตอบคำถามหลังนำเสนอ"),
        ("หลักการตีความ", "รายงานผลเป็น trade-off; ไม่ประกาศผู้ชนะทุก metric"),
        ("สถานะ", "พร้อมให้อาจารย์ตรวจ; human-expert validation ยังรอก่อน submission"),
    )
    for row, values in zip(table.rows, rows):
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            set_cell_shading(cell, LIGHT_BLUE if idx == 0 else "FFFFFF")
            p = cell.paragraphs[0]
            run = p.add_run(value)
            set_run_font(run, 14, bold=idx == 0, color=TEAL if idx == 0 else None)

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("คำเตือน: Radar charts เป็น descriptive visualization ไม่ใช่ composite inferential score")
    set_run_font(r, 13, bold=True, color=ORANGE)
    shade_paragraph(p, LIGHT_ORANGE)


def build() -> Path:
    if not TALK.exists() or not QA.exists():
        raise SystemExit("Missing talk or Q&A Markdown source")
    document = Document()
    configure_document(document)
    add_cover(document)
    document.add_page_break()
    document.add_heading("ส่วนที่ 1  บทพูดนำเสนอ 10 นาที", level=1)
    add_markdown(document, TALK)
    document.add_page_break()
    document.add_heading("ส่วนที่ 2  คำถาม–คำตอบสำหรับอาจารย์", level=1)
    add_markdown(document, QA)

    document.core_properties.title = "ConveyorFlow Advisor Talk and Q&A"
    document.core_properties.subject = "Advisor presentation support document"
    document.core_properties.author = "Sakan Punyanon"
    document.core_properties.keywords = "ConveyorFlow, LLM agents, advisor, talk, Q&A"
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
