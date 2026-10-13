"""Create Rev10 from Rev9 with a checked, already-pushed evidence commit."""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document


CURRENT = Path(__file__).resolve().parent.parent / "CURRENT_MANUSCRIPTS"
PREFIXES = (
    ("ConveyorFlow_IEEE_Manuscript_v2_7_Exploratory_Draft", False),
    ("ConveyorFlow_AJSTR_Blinded_v2_7_Exploratory_Draft", True),
    ("ConveyorFlow_AJSTR_Unblinded_v2_7_Exploratory_Draft", False),
)
OLD = "691bc9e1397b45f9948069d4b6b44b0e48d8e826"


def windows_safe(path: Path) -> str:
    value = str(path.resolve())
    return "\\\\?\\" + value if len(value) >= 240 and not value.startswith("\\\\?\\") else value


def media(path: Path):
    with ZipFile(windows_safe(path)) as archive:
        return {name: archive.read(name) for name in archive.namelist()
                if name.startswith("word/media/")}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("commit", help="40-character hash already pushed to GitHub")
    args = parser.parse_args()
    if not re.fullmatch(r"[0-9a-f]{40}", args.commit):
        raise ValueError("A complete Git commit hash is required")
    replacement = args.commit
    for stem, blind in PREFIXES:
        source = CURRENT / (stem + "_Rev9_ReferencesVerified.docx")
        target = CURRENT / (stem + "_Rev10_ReferencesVerified_GitPinned.docx")
        if target.exists():
            raise FileExistsError(target)
        doc = Document(windows_safe(source))
        changed = 0
        if not blind:
            for paragraph in doc.paragraphs:
                if OLD in paragraph.text:
                    paragraph.text = paragraph.text.replace(OLD, replacement)
                    paragraph.text = paragraph.text.replace(
                        "(frozen commit 691bc9e)",
                        f"(frozen commit {replacement[:7]})")
                    changed += 1
            if changed != 1:
                raise ValueError(f"Expected one data/code paragraph in {source.name}; got {changed}")
        else:
            full_text = "\n".join(p.text for p in doc.paragraphs)
            if OLD in full_text or replacement in full_text or "github.com/sakanarm" in full_text:
                raise ValueError("Blinded manuscript discloses a repository identifier")
        before = (len(doc.inline_shapes), len(doc.tables), len(doc.element.body.xpath(".//m:oMath")))
        doc.save(windows_safe(target))
        after = Document(windows_safe(target))
        if before != (len(after.inline_shapes), len(after.tables), len(after.element.body.xpath(".//m:oMath"))):
            raise ValueError("Document structure changed")
        if media(source) != media(target):
            raise ValueError("Embedded image changed")
        print(target.name, "changed availability paragraphs", changed)


if __name__ == "__main__":
    main()
