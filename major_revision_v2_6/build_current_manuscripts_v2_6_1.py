"""Versioned v2.6.1 manuscript drafts from the latest CURRENT_MANUSCRIPTS.

The original six-block result and separate 11-block cohort remain distinct.
This file is a fallback authoring route when Microsoft Word COM is unavailable;
the outputs require Word visual review before journal submission.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from build_bounded_manuscripts_v1 import (
    ABSTRACT, CONCLUSION_ADD, LIMIT, METHOD, RESULT,
    find, insert_before, table_before, windows_safe,
)


V2 = Path(__file__).resolve().parents[1]
CURRENT = V2 / "CURRENT_MANUSCRIPTS"
SOURCES = [
    ("ConveyorFlow_IEEE_Manuscript_v2_3_Main_Rev8.docx",
     "ConveyorFlow_IEEE_Manuscript_v2_6_1_Current.docx", False, False),
    ("ConveyorFlow_AJSTR_v2_3_Blinded_Manuscript_Main_Rev6.docx",
     "ConveyorFlow_AJSTR_Blinded_v2_6_1_Current.docx", True, True),
    ("ConveyorFlow_AJSTR_v2_3_Unblinded_Manuscript_Main_Rev6.docx",
     "ConveyorFlow_AJSTR_Unblinded_v2_6_1_Current.docx", True, False),
]

PILOT = (
    "A later instrument diagnostic used one previously exposed FastAPI repair "
    "case, not a new first-attempt sample or allocation comparison. Exploratory "
    "CF-Fit and central requests had each returned an empty length-limited "
    "response at a 4,096-token output ceiling (0/1 verified per arm). A single "
    "frozen, no-retry call to the same deployment with a 32,768-token ceiling "
    "returned a complete exact-edit patch and passed the visible test, three "
    "public regression tests, and fresh replay of both checks. The cap change "
    "was made after seeing the earlier failures. This result is an exposed-case "
    "instrument check, excluded from all cohort denominators; it estimates no "
    "allocation effect."
)

INTERPRETATION = (
    "The 11-block executable cohort is separate from the earlier six-block "
    "experiment and is not an independent replication of its datasets or "
    "repositories. CF-Fit verified more jobs than the rule-matched central "
    "arm, but Static-Owners verified the most in the 11-block cohort. This is "
    "the intended throughput, time, utilization, and cost trade-off question, "
    "not a universal policy ranking. The live comparison still includes "
    "provider timing and model-output variation; only the small matched fault "
    "fixture isolates the allocation decision process under controlled outcomes."
)


def update(doc: Document, ajstr: bool, blind: bool) -> None:
    abstract = next(p for p in doc.paragraphs if p.text.startswith("Abstract"))
    abstract.text = ("Abstract: " if ajstr else "Abstract\u2014") + ABSTRACT + (
        " The earlier six-block executable study is reported separately; "
        "the cohorts are not pooled."
    )
    if not ajstr:
        status = next(p for p in doc.paragraphs if p.text.startswith("Research revision 2.3"))
        status.text = (
            "Research revision 2.6.1, 13 October 2026. The earlier six-block "
            "and separate 11-block executable cohorts are independently audited "
            "and not pooled. Matched-fault results are fixture-level evidence. "
            "Human difficulty validation, broader source replication, complete "
            "billing, coauthor review, and final journal-format review remain open."
        )
    method_anchor = find(doc, "3. Results and Discussion" if ajstr else "VI. RESULTS")
    insert_before(method_anchor,
                  "2.5. Separate executable cohort and matched fault control"
                  if ajstr else
                  "J. Separate Executable Cohort and Matched Fault Control",
                  "Heading 2")
    insert_before(method_anchor,
                  "The earlier six-block ecological main and the later 11-block "
                  "executable cohort are separate experiments; they are not "
                  "combined or counted as 17 independent blocks.")
    for paragraph in METHOD:
        insert_before(method_anchor, paragraph)
    result_anchor = find(doc, "3.12. Interpretation" if ajstr else "VII. DISCUSSION")
    insert_before(result_anchor,
                  "3.12. Separate 11-block executable cohort and fault results"
                  if ajstr else
                  "O. Separate 11-Block Executable Cohort and Fault Results",
                  "Heading 2")
    insert_before(result_anchor,
                  "Table 12. Separate 11-block executable cohort"
                  if ajstr else
                  "TABLE XIII. SEPARATE 11-BLOCK EXECUTABLE COHORT")
    table_before(result_anchor, [
        ("Arm", "Verified", "Full ML", "Repo repair", "Jobs/hour", "Unknown-cost calls"),
        ("CF-Fit", "17/22", "9/11", "8/11", "1.545", "1"),
        ("Central-Rule-Matched", "14/22", "8/11", "6/11", "1.273", "1"),
        ("Static-Owners", "18/22", "9/11", "9/11", "1.636", "3"),
    ], doc)
    for paragraph in RESULT[:2]:
        insert_before(result_anchor, paragraph)
    insert_before(result_anchor,
                  "Table 13. Matched decision-layer fault outcomes during the fault"
                  if ajstr else
                  "TABLE XIV. MATCHED DECISION-LAYER FAULT OUTCOMES")
    table_before(result_anchor, [
        ("Fault scenario", "CF-Fit", "Central-Rule-Matched"),
        ("No fault", "3/3", "3/3"),
        ("Coordinator-only outage", "3/3", "0/3"),
        ("Shared belt/claim-store outage", "0/3", "0/3"),
        ("Agent A2 outage", "2/3", "2/3"),
    ], doc)
    insert_before(result_anchor, RESULT[2])
    insert_before(result_anchor, PILOT)
    if ajstr:
        find(doc, "3.12. Interpretation").text = "3.13. Interpretation"
        find(doc, "3.13. Limitations").text = "3.14. Limitations"
        discussion_anchor = find(doc, "3.14. Limitations")
        limit_anchor = find(doc, "4. Conclusions")
        conclusion_anchor = find(doc, "Data and code availability")
    else:
        discussion_anchor = find(doc, "VIII. THREATS TO VALIDITY")
        limit_anchor = find(doc, "IX. CONCLUSION")
        conclusion_anchor = find(doc, "DATA AND CODE AVAILABILITY")
    insert_before(discussion_anchor, INTERPRETATION)
    insert_before(limit_anchor, LIMIT + " " +
                  "The exposed FastAPI output-budget pilot is post-outcome "
                  "instrument calibration, not prospective cohort evidence.")
    insert_before(conclusion_anchor, CONCLUSION_ADD)
    doc.core_properties.title = "ConveyorFlow decentralized task allocation manuscript"
    doc.core_properties.subject = "Separate executable cohort and matched fault evidence"
    if blind:
        doc.core_properties.author = ""
        doc.core_properties.last_modified_by = ""


def main() -> None:
    for source_name, output_name, ajstr, blind in SOURCES:
        source, output = CURRENT / source_name, CURRENT / output_name
        if output.exists():
            raise FileExistsError(f"Refuse to overwrite {output}")
        doc = Document(windows_safe(source))
        before = (len(doc.inline_shapes), len(doc.element.body.xpath('.//m:oMath')))
        update(doc, ajstr, blind)
        after = (len(doc.inline_shapes), len(doc.element.body.xpath('.//m:oMath')))
        if after != before:
            raise ValueError(f"Figure/equation preservation failed for {source_name}")
        doc.save(windows_safe(output))
        reopened = Document(windows_safe(output))
        print(f"{output_name}: {len(reopened.inline_shapes)} figures, "
              f"{len(reopened.tables)} tables, {after[1]} native equations")


if __name__ == "__main__":
    main()
