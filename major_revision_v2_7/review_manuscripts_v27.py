"""Read-only cross-section review of the three current manuscript drafts."""

from __future__ import annotations

import json
import re

from docx import Document

from build_manuscript_v27_drafts import CURRENT, SOURCES, windows_safe


def inspect(name, ajstr):
    doc = Document(windows_safe(CURRENT / name))
    paragraphs = doc.paragraphs
    body = [p.text for p in paragraphs]
    figures = {}
    tables = {}
    for index, value in enumerate(body):
        pattern = (r"^Figure (\d+)\." if ajstr else r"^Fig\. (\d+)\.")
        match = re.match(pattern, value.strip())
        if match:
            figures[int(match.group(1))] = index
        pattern = (r"^Table (\d+)\." if ajstr else r"^TABLE ([IVX]+)\.")
        match = re.match(pattern, value.strip())
        if match:
            tables[match.group(1)] = index
    figure_uncited = []
    for number, index in figures.items():
        pattern = re.compile(rf"\b(?:Figure|Fig\.)\s*{number}\b")
        grouped = any(f"Figures {number} and " in value for value in body) or any(
            f"Figures {first} and {number}" in value
            for first in range(1, number) for value in body)
        if not grouped and not any(pattern.search(value) for i, value in enumerate(body)
                                   if i != index):
            figure_uncited.append(number)
    table_uncited = []
    for number, index in tables.items():
        pattern = re.compile(rf"\bTable\s*{number}\b", re.IGNORECASE)
        if not any(pattern.search(value) for i, value in enumerate(body) if i != index):
            table_uncited.append(number)
    blank_heading_images = [i for i, p in enumerate(paragraphs)
                            if p.style.name.startswith("Heading")
                            and not p.text.strip()
                            and p._p.xpath(".//w:drawing")]
    abstract = next((value for value in body if value.startswith("Abstract")), "")
    reference_heading = next(i for i, value in enumerate(body)
                             if value in ("References", "REFERENCES"))
    first_citations = []
    for value in body[:reference_heading]:
        for group in re.findall(r"\[([0-9,\s\-–]+)\]", value):
            for part in re.split(r",\s*", group):
                endpoints = re.split(r"\s*[\-–]\s*", part.strip())
                values = (range(int(endpoints[0]), int(endpoints[1]) + 1)
                          if len(endpoints) == 2 else [int(endpoints[0])])
                for number in values:
                    if number not in first_citations:
                        first_citations.append(number)
    return {
        "file": name,
        "paragraphs": len(paragraphs),
        "abstract_words": len(abstract.split()),
        "figure_captions": figures,
        "table_captions": tables,
        "figures_without_body_citation": figure_uncited,
        "tables_without_body_citation": table_uncited,
        "blank_heading_image_paragraphs": blank_heading_images,
        "reference_first_appearance_order": first_citations,
    }


def main():
    rows = [inspect(name, ajstr) for _, name, ajstr, _ in SOURCES]
    blind = Document(windows_safe(CURRENT / SOURCES[1][1]))
    unblind = Document(windows_safe(CURRENT / SOURCES[2][1]))
    left = [p.text for p in blind.paragraphs]
    right = [p.text for p in unblind.paragraphs]
    left = left[left.index("1. Introduction"):]
    right = right[right.index("1. Introduction"):]
    if len(left) != len(right):
        raise ValueError("AJSTR blind/unblind section structure differs")
    differences = [{"relative_paragraph": i, "blind_start": a[:80],
                    "unblind_start": b[:80]}
                   for i, (a, b) in enumerate(zip(left, right)) if a != b]
    print(json.dumps({"documents": rows,
                      "ajstr_text_differences_after_introduction": differences},
                     indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
