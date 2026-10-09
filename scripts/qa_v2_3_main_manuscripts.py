"""Read-only consistency check for the current IEEE and AJSTR main manuscripts.

This supplements, but does not replace, page-by-page review in Microsoft Word.
"""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1] / "CURRENT_MANUSCRIPTS"
FILES = {
    "IEEE": ROOT / "ConveyorFlow_IEEE_Manuscript_v2_3_Main_Rev7.docx",
    "AJSTR_blind": ROOT / "ConveyorFlow_AJSTR_v2_3_Blinded_Manuscript_Main_Rev5.docx",
    "AJSTR_unblind": ROOT / "ConveyorFlow_AJSTR_v2_3_Unblinded_Manuscript_Main_Rev5.docx",
}


def body_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def check(name: str, path: Path) -> dict[str, object]:
    require(path.is_file(), f"Missing manuscript: {path}")
    document = Document(path)
    text = body_text(document)
    normal = re.sub(r"\s+", " ", text)
    for count in ("11/18", "10/18", "12/18"):
        require(count in normal, f"{name}: missing verified-job count {count}")
    require("05_R1" in normal, f"{name}: missing interrupted-block replacement")
    require("Matplotlib" in normal, f"{name}: missing one verified repository repair attribution")
    require("CF-Fit" in normal and "Central" in normal and "Static" in normal,
            f"{name}: missing an allocation arm")
    require("3600" in normal or "3,600" in normal,
            f"{name}: missing fixed observation horizon")
    require("provider" in normal.lower() and "cost" in normal.lower(),
            f"{name}: missing provider-cost context")
    expected_figures = 15 if name == "IEEE" else 16
    require(len(document.inline_shapes) == expected_figures,
            f"{name}: expected {expected_figures} embedded figures")
    require(len(document.element.xpath(".//m:oMath")) == 13,
            f"{name}: expected 13 native Word equations")
    with zipfile.ZipFile(path) as package:
        if name == "AJSTR_blind":
            package_text = "\n".join(
                package.read(member).decode("utf-8", "replace")
                for member in package.namelist()
                if member.endswith((".xml", ".rels"))
            ).lower()
            for identity in ("sakan", "punyanon", "chaiyaporn", "aomduan",
                             "dpu.ac.th", "68140010", "sakanarm"):
                require(identity not in package_text,
                        f"{name}: author identity in package: {identity}")
        if name == "AJSTR_unblind":
            for identity in ("Sakan Punyanon", "Chaiyaporn", "Aomduan"):
                require(identity in text, f"{name}: missing author {identity}")
    return {
        "format": name,
        "file": str(path),
        "figures": len(document.inline_shapes),
        "tables": len(document.tables),
        "word_equations": 13,
        "six_block_counts_present": True,
        "replacement_disclosed": True,
        "visual_qa": "performed separately in Microsoft Word",
    }


def main() -> None:
    reports = [check(name, path) for name, path in FILES.items()]
    print(json.dumps(reports, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
