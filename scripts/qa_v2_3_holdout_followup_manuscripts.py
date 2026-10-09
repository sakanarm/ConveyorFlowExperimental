"""Read-only checks for the Word-authored post-main holdout-repair revisions."""

from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1] / "CURRENT_MANUSCRIPTS"
SPECS = (
    ("IEEE", "ConveyorFlow_IEEE_Manuscript_v2_3_Main_Rev7.docx", "ConveyorFlow_IEEE_Manuscript_v2_3_Main_Rev8.docx", 15, 12),
    ("AJSTR_blind", "ConveyorFlow_AJSTR_v2_3_Blinded_Manuscript_Main_Rev5.docx", "ConveyorFlow_AJSTR_v2_3_Blinded_Manuscript_Main_Rev6.docx", 16, 11),
    ("AJSTR_unblind", "ConveyorFlow_AJSTR_v2_3_Unblinded_Manuscript_Main_Rev5.docx", "ConveyorFlow_AJSTR_v2_3_Unblinded_Manuscript_Main_Rev6.docx", 16, 11),
)


def text(document: Document) -> str:
    return "\n".join(p.text for p in document.paragraphs)


def main() -> None:
    records = []
    for name, old_name, new_name, figures, tables in SPECS:
        source = ROOT / old_name
        output = ROOT / new_name
        before = Document(source)
        after = Document(output)
        old_text = text(before)
        new_text = text(after)
        assert len(after.inline_shapes) == len(before.inline_shapes) == figures, name
        assert len(after.tables) == len(before.tables) == tables, name
        assert len(after.element.xpath(".//m:oMath")) == len(before.element.xpath(".//m:oMath")) == 13, name
        assert "15/15 started calls" in new_text, name
        assert "7 verified repairs" in new_text, name
        assert "Six cases were prespecified" in new_text, name
        assert "excluded before any model call" in new_text, name
        assert "not pooled" in new_text, name
        assert "CF-Fit and Central-Fit verified no repository repair" in new_text, name
        assert "11/18" in new_text and "10/18" in new_text and "12/18" in new_text, name
        assert "15/15 started calls" not in old_text, name
        if name == "AJSTR_blind":
            with zipfile.ZipFile(output) as package:
                all_xml = "\n".join(
                    package.read(member).decode("utf-8", "replace")
                    for member in package.namelist()
                    if member.endswith((".xml", ".rels"))
                ).lower()
            for identity in ("sakan", "punyanon", "chaiyaporn", "aomduan", "dpu.ac.th", "68140010", "sakanarm"):
                assert identity not in all_xml, f"Blind identity leak: {identity}"
        records.append({
            "format": name,
            "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "output_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
            "figures": figures,
            "tables": tables,
            "word_equations": 13,
            "followup_separate_from_main": True,
            "visual_qa_required": True,
        })
    print(json.dumps(records, indent=2))


if __name__ == "__main__":
    main()
