"""Read-only structural/identity check; Word page QA is a separate gate."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import zipfile

from docx import Document
from docx.oxml.ns import qn

from build_manuscript_v27_drafts import (CURRENT, FIGURE5, FROZEN_URL,
                                         SOURCES, windows_safe)


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def media(path):
    with zipfile.ZipFile(windows_safe(path)) as archive:
        return {name: hashlib.sha256(archive.read(name)).hexdigest()
                for name in archive.namelist() if name.startswith("word/media/")}


def inline_image_hashes(doc):
    values = []
    for paragraph in doc.paragraphs:
        for blip in paragraph._p.xpath(".//a:blip"):
            rid = blip.get(qn("r:embed"))
            values.append(hashlib.sha256(doc.part.related_parts[rid].blob).hexdigest())
    return values


def audit():
    rows = []
    for source_name, target_name, ajstr, blind in SOURCES:
        source = CURRENT / source_name
        target = CURRENT / target_name
        old = Document(windows_safe(source))
        new = Document(windows_safe(target))
        text = "\n".join(paragraph.text for paragraph in new.paragraphs)
        required = ("Bounded Scrapy", "1,800-s horizon", "0.024945",
                    "0.026529", "0.023978", "Future work",
                    "order in which READY tasks", "one repository project",
                    "not verified billed costs")
        missing = [item for item in required if item.lower() not in text.lower()]
        if missing:
            raise ValueError(f"{target_name}: missing claims {missing}")
        old_images = inline_image_hashes(old)
        new_images = inline_image_hashes(new)
        changed = [i for i, (a, b) in enumerate(zip(old_images, new_images))
                   if a != b]
        if (len(old_images) != len(new_images) or len(changed) != 1 or
                new_images[changed[0]] != sha(FIGURE5)):
            raise ValueError(f"{target_name}: unexpected figure drift {changed}")
        old_media, new_media = media(source), media(target)
        if (any(new_media.get(key) != value for key, value in old_media.items()) or
                len(new_media) != len(old_media) + 1):
            raise ValueError(f"{target_name}: unexpected media-package drift")
        before = (len(old.inline_shapes), len(old.element.body.xpath('.//m:oMath')),
                  len(old.tables))
        after = (len(new.inline_shapes), len(new.element.body.xpath('.//m:oMath')),
                 len(new.tables))
        if after != (before[0], before[1], before[2] + 1):
            raise ValueError(f"{target_name}: figure/equation/table drift {before} -> {after}")
        if ajstr:
            for heading in ("3.13. Bounded Scrapy repository extension",
                            "3.14. Interpretation", "3.15. Limitations"):
                if text.count(heading) != 1:
                    raise ValueError(f"{target_name}: heading drift {heading}")
        if blind:
            if FROZEN_URL in text or "editorial office separately" not in text:
                raise ValueError(f"{target_name}: blind availability drift")
        elif text.count(FROZEN_URL) != 1:
            raise ValueError(f"{target_name}: frozen repository link missing or duplicated")
        if blind:
            with zipfile.ZipFile(windows_safe(target)) as archive:
                xml = b"\n".join(archive.read(name) for name in archive.namelist()
                                 if name.endswith((".xml", ".rels"))).lower()
            for token in (b"sakan", b"punyanon", b"chaiyaporn", b"aomduan",
                          b"68140010", b"github.com/sakanarm"):
                if token in xml:
                    raise ValueError(f"Blind identifier leaked: {token!r}")
        rows.append({"file": target_name, "sha256": sha(target),
                     "figures": after[0], "equations": after[1],
                     "tables": after[2], "other_figures_preserved": True,
                     "figure5_corrected": True,
                     "blind_identity_check": "passed" if blind else "not_applicable",
                     "structural_check": "passed"})
    return rows


if __name__ == "__main__":
    print(json.dumps(audit(), indent=2))
