"""Audit reference numbering, citation use, and figure preservation in Rev9."""

from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from docx import Document


CURRENT = Path(__file__).resolve().parent.parent / "CURRENT_MANUSCRIPTS"
STEMS = (
    "ConveyorFlow_IEEE_Manuscript_v2_7_Exploratory_Draft",
    "ConveyorFlow_AJSTR_Blinded_v2_7_Exploratory_Draft",
    "ConveyorFlow_AJSTR_Unblinded_v2_7_Exploratory_Draft",
)


def citation_numbers(text: str):
    text = re.sub(r"\[(\d+)\]\s*[–-]\s*\[(\d+)\]", r"[\1–\2]", text)
    for group in re.findall(r"\[([0-9,\s–-]+)\]", text):
        for part in group.split(","):
            ends = re.split(r"\s*[–-]\s*", part.strip())
            if len(ends) == 1:
                yield int(ends[0])
            elif len(ends) == 2:
                yield from range(int(ends[0]), int(ends[1]) + 1)


def media(path: Path):
    with ZipFile(path) as archive:
        return Counter(
            (name.rsplit(".", 1)[-1], __import__("hashlib").sha256(archive.read(name)).hexdigest())
            for name in archive.namelist() if name.startswith("word/media/")
        )


def inspect(stem: str):
    earlier = CURRENT / (stem + "_Rev7.docx")
    current = CURRENT / (stem + "_Rev9_ReferencesVerified.docx")
    doc = Document(str(current))
    older = Document(str(earlier))
    paragraphs = [p.text for p in doc.paragraphs]
    heading = "References" if "AJSTR" in stem else "REFERENCES"
    boundary = paragraphs.index(heading)
    refs = [(int(match.group(1)), value) for value in paragraphs[boundary + 1:]
            if (match := re.match(r"^\[(\d+)\]\s+", value))]
    citations = []
    for value in paragraphs[:boundary]:
        citations.extend(citation_numbers(value))
    first = list(dict.fromkeys(citations))
    expected = list(range(1, 27))
    if [number for number, _ in refs] != expected:
        raise ValueError(f"Reference numbering is incomplete or unordered in {stem}")
    if first != expected:
        raise ValueError(f"First citation order is wrong in {stem}: {first}")
    if any(number not in citations for number in expected):
        raise ValueError(f"Uncited entry in {stem}")
    if len(doc.inline_shapes) != len(older.inline_shapes):
        raise ValueError(f"Figure count changed in {stem}")
    if len(doc.tables) != len(older.tables):
        raise ValueError(f"Table count changed in {stem}")
    if len(doc.element.body.xpath(".//m:oMath")) != len(older.element.body.xpath(".//m:oMath")):
        raise ValueError(f"Equation count changed in {stem}")
    if media(current) != media(earlier):
        raise ValueError(f"An embedded image changed in {stem}")
    if "AJSTR" in stem:
        for number, value in refs:
            if "et al." in value and len(value.split(";")) < 11:
                raise ValueError(f"Too few named authors in reference {number}")
        after_punctuation = [(index, value[:100]) for index, value in enumerate(paragraphs[:boundary])
                             if re.search(r"[.!?]\s+\[\d+\]", value)]
    else:
        after_punctuation = []
    return {
        "file": current.name,
        "references": len(refs),
        "first_citation_order": first,
        "figures": len(doc.inline_shapes),
        "tables": len(doc.tables),
        "equations": len(doc.element.body.xpath(".//m:oMath")),
        "embedded_images_unchanged": True,
        "words_in_paragraphs_and_tables": sum(len(value.split()) for value in paragraphs)
            + sum(len(cell.text.split()) for table in doc.tables for row in table.rows for cell in row.cells),
        "citations_after_sentence_punctuation": after_punctuation,
    }


def main():
    print(json.dumps([inspect(stem) for stem in STEMS], indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
